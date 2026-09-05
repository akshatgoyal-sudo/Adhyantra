from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from fpdf.enums import MethodReturnValue


PROJECT_ROOT = Path(__file__).resolve().parents[2]
FONT_ROOT = PROJECT_ROOT / "frontend" / "public" / "fonts"
INSTRUMENT_FONT = FONT_ROOT / "instrument-sans-latin-variable.woff2"
SOURCE_SERIF_FONT = FONT_ROOT / "source-serif-4-latin-variable.woff2"
DEVANAGARI_FONT = FONT_ROOT / "noto-sans-devanagari-variable.woff2"
DOCUMENT_EXPORT_FORMATS = frozenset({"pdf_export", "docx_export"})
DEVANAGARI_PATTERN = re.compile(r"[\u0900-\u097f]")


class LessonDocumentGenerationError(RuntimeError):
    pass


@dataclass(frozen=True)
class LessonNotesSection:
    title: str
    summary: str | None
    bullets: tuple[str, ...]
    examples: tuple[str, ...]
    remember_points: tuple[str, ...]
    revision_cues: tuple[str, ...]


@dataclass(frozen=True)
class LessonNotesContent:
    title: str
    exam: str
    subject: str
    topic: str
    teaching_mode: str
    lesson_mode: str
    generated_at: str
    simple_explanation: str
    detailed_explanation: str | None
    sections: tuple[LessonNotesSection, ...]
    questions: tuple[str, ...]


def _text(value: Any) -> str:
    return str(value or "").strip()


def _text_items(value: Any) -> tuple[str, ...]:
    if not isinstance(value, list):
        return ()
    return tuple(text for item in value if (text := _text(item)))


def _normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value)).casefold()


def _unique_value(value: Any, seen: set[str]) -> str | None:
    text = _text(value)
    normalized = _normalize(text)
    if not normalized or normalized in seen:
        return None
    seen.add(normalized)
    return text


def _unique_items(value: Any, seen: set[str]) -> tuple[str, ...]:
    items: list[str] = []
    for candidate in value if isinstance(value, list) else []:
        text = _unique_value(candidate, seen)
        if text:
            items.append(text)
    return tuple(items)


def _document_sections(lesson: dict[str, Any], seen: set[str]) -> tuple[LessonNotesSection, ...]:
    source = lesson.get("sections") if isinstance(lesson.get("sections"), list) else []
    if not source:
        source = [
            {"title": "Key points", "bullets": lesson.get("key_points") or []},
            {"title": "Examples", "examples": lesson.get("examples") or []},
            {"title": "Exam relevance", "summary": lesson.get("exam_relevance")},
            {"title": "Common traps", "bullets": lesson.get("common_traps") or []},
        ]

    sections: list[LessonNotesSection] = []
    for raw in source:
        if not isinstance(raw, dict):
            continue
        title = _text(raw.get("title"))
        summary = _unique_value(raw.get("summary"), seen)
        bullets = _unique_items(raw.get("bullets"), seen)
        examples = _unique_items(raw.get("examples"), seen)
        remember_points = _unique_items(raw.get("remember_points"), seen)
        revision_cues = _unique_items(raw.get("revision_cues"), seen)
        if not title or not any((summary, bullets, examples, remember_points, revision_cues)):
            continue
        sections.append(
            LessonNotesSection(
                title=title,
                summary=summary,
                bullets=bullets,
                examples=examples,
                remember_points=remember_points,
                revision_cues=revision_cues,
            )
        )
    return tuple(sections)


def build_lesson_notes_content(lesson: dict[str, Any], generated_at: str) -> LessonNotesContent:
    simple = _text(lesson.get("simple_explanation"))
    detailed = _text(lesson.get("detailed_explanation"))
    seen = {_normalize(simple)} if simple else set()
    if detailed and _normalize(detailed) in seen:
        detailed = ""
    elif detailed:
        seen.add(_normalize(detailed))
    return LessonNotesContent(
        title=_text(lesson.get("topic")) or "Adhyantra Lesson Notes",
        exam=_text(lesson.get("exam")).upper(),
        subject=_text(lesson.get("subject")).replace("_", " ").title(),
        topic=_text(lesson.get("topic")) or "Lesson",
        teaching_mode=_text(lesson.get("teaching_mode")).replace("_", " ").title(),
        lesson_mode=_text(lesson.get("lesson_mode")).replace("_", " ").title(),
        generated_at=generated_at,
        simple_explanation=simple,
        detailed_explanation=detailed or None,
        sections=_document_sections(lesson, seen),
        questions=_text_items(lesson.get("practice_questions"))[:5],
    )


def _formatted_date(value: str) -> str:
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).strftime("%d %B %Y")
    except (TypeError, ValueError):
        return value[:10] if value else ""


def _ensure_pdf_space(pdf: FPDF, height: float) -> None:
    if pdf.will_page_break(height):
        pdf.add_page()


def _pdf_font(pdf: FPDF, family: str, size: float, color: tuple[int, int, int] = (26, 39, 46)) -> None:
    pdf.set_font(family, size=size)
    pdf.set_text_color(*color)


def _pdf_actual_text(value: str) -> str:
    return "FEFF" + value.encode("utf-16-be").hex().upper()


def _pdf_multiline(pdf: FPDF, width: float, height: float, text: str, *, indent: float = 0) -> None:
    lines = pdf.multi_cell(width, height, text, dry_run=True, output=MethodReturnValue.LINES)
    for line in lines:
        _ensure_pdf_space(pdf, height)
        pdf.set_x(pdf.l_margin + indent)
        if DEVANAGARI_PATTERN.search(line):
            pdf._out(f"/Span <</ActualText <{_pdf_actual_text(line)}>>> BDC")
        pdf.multi_cell(width, height, line, new_x="LMARGIN", new_y="NEXT")
        if DEVANAGARI_PATTERN.search(line):
            pdf._out("EMC")


def _pdf_paragraph(pdf: FPDF, text: str, *, size: float = 11.2, indent: float = 0) -> None:
    if not text:
        return
    _pdf_font(pdf, "SourceSerif", size)
    pdf.set_x(pdf.l_margin + indent)
    _pdf_multiline(pdf, pdf.epw - indent, 6.4, text, indent=indent)
    pdf.ln(2.2)


def _pdf_heading(pdf: FPDF, text: str, *, level: int = 1) -> None:
    _ensure_pdf_space(pdf, 22 if level == 1 else 16)
    pdf.ln(2 if level == 1 else 1)
    _pdf_font(pdf, "Instrument", 15.5 if level == 1 else 11.5, (20, 52, 61) if level == 1 else (15, 118, 110))
    _pdf_multiline(pdf, pdf.epw, 8 if level == 1 else 6.5, text)
    pdf.ln(1.5)


def _pdf_list(pdf: FPDF, items: Iterable[str], *, numbered: bool = False) -> None:
    for index, item in enumerate(items, start=1):
        prefix = f"{index}. " if numbered else "- "
        _pdf_font(pdf, "SourceSerif", 10.8)
        lines = pdf.multi_cell(pdf.epw - 3, 6.2, f"{prefix}{item}", dry_run=True, output=MethodReturnValue.LINES)
        estimated_height = len(lines) * 6.2 + 1
        if estimated_height <= pdf.h - pdf.t_margin - pdf.b_margin:
            _ensure_pdf_space(pdf, estimated_height)
        pdf.set_x(pdf.l_margin + 3)
        _pdf_multiline(pdf, pdf.epw - 3, 6.2, f"{prefix}{item}", indent=3)
        pdf.ln(1)


def build_pdf_lesson_notes(lesson: dict[str, Any], generated_at: str) -> bytes:
    try:
        notes = build_lesson_notes_content(lesson, generated_at)
        pdf = FPDF(format="A4")
        pdf.set_margins(18, 17, 18)
        pdf.set_auto_page_break(True, 17)
        pdf.add_font("Instrument", fname=str(INSTRUMENT_FONT))
        pdf.add_font("SourceSerif", fname=str(SOURCE_SERIF_FONT))
        pdf.add_font("NotoDevanagari", fname=str(DEVANAGARI_FONT))
        pdf.set_fallback_fonts(["NotoDevanagari"])
        pdf.set_text_shaping(True)
        pdf.set_title(notes.title)
        pdf.set_author("Adhyantra")
        pdf.set_creator("Adhyantra")
        pdf.add_page()

        _pdf_font(pdf, "Instrument", 12, (15, 143, 131))
        pdf.cell(0, 7, "Adhyantra")
        pdf.ln(11)
        _pdf_font(pdf, "Instrument", 24, (16, 42, 53))
        _pdf_multiline(pdf, pdf.epw, 11, notes.title)
        pdf.ln(3)
        metadata = (
            ("Exam", notes.exam),
            ("Subject", notes.subject),
            ("Topic", notes.topic),
            ("Teaching approach", notes.teaching_mode),
            ("Lesson mode", notes.lesson_mode),
            ("Exported", _formatted_date(notes.generated_at)),
        )
        for label, value in metadata:
            if not value:
                continue
            _pdf_font(pdf, "Instrument", 8.8, (64, 81, 90))
            pdf.cell(35, 5.5, label)
            _pdf_font(pdf, "SourceSerif", 9.7)
            _pdf_multiline(pdf, pdf.epw - 35, 5.5, value, indent=35)
        pdf.ln(4)

        _pdf_heading(pdf, "Core explanation")
        if notes.simple_explanation:
            _pdf_font(pdf, "SourceSerif", 12)
            callout_lines = pdf.multi_cell(pdf.epw - 4, 6.4, notes.simple_explanation, dry_run=True, output=MethodReturnValue.LINES)
            callout_height = len(callout_lines) * 6.4 + 2.2
            if callout_height <= pdf.h - pdf.t_margin - pdf.b_margin:
                _ensure_pdf_space(pdf, callout_height)
                start_y = pdf.get_y()
                _pdf_paragraph(pdf, notes.simple_explanation, size=12, indent=4)
                end_y = pdf.get_y() - 2
                pdf.set_draw_color(15, 143, 131)
                pdf.set_line_width(0.7)
                pdf.line(pdf.l_margin, start_y, pdf.l_margin, end_y)
            else:
                _pdf_paragraph(pdf, notes.simple_explanation, size=12)
        if notes.detailed_explanation:
            _pdf_paragraph(pdf, notes.detailed_explanation)

        for section in notes.sections:
            _pdf_heading(pdf, section.title)
            if section.summary:
                _pdf_paragraph(pdf, section.summary)
            if section.bullets:
                _pdf_list(pdf, section.bullets)
            if section.examples:
                _pdf_heading(pdf, "Examples", level=2)
                _pdf_list(pdf, section.examples)
            if section.remember_points:
                _pdf_heading(pdf, "Remember", level=2)
                _pdf_list(pdf, section.remember_points)
            if section.revision_cues:
                _pdf_heading(pdf, "Revision cues", level=2)
                _pdf_list(pdf, section.revision_cues)

        if notes.questions:
            _pdf_heading(pdf, "Check your understanding")
            _pdf_list(pdf, notes.questions, numbered=True)

        result = bytes(pdf.output())
        if not result.startswith(b"%PDF-"):
            raise LessonDocumentGenerationError("The generated PDF did not have a valid signature.")
        return result
    except LessonDocumentGenerationError:
        raise
    except Exception as exc:
        raise LessonDocumentGenerationError("The PDF study notes could not be generated.") from exc


def _set_run_font(run, font_name: str, size: float | None = None) -> None:
    run.font.name = font_name
    if size is not None:
        run.font.size = Pt(size)
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)


def _docx_text_font(text: str) -> str:
    return "Nirmala UI" if DEVANAGARI_PATTERN.search(text) else "Georgia"


def _add_docx_paragraph(document: Document, text: str, *, style: str | None = None, bold_lead: str | None = None):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.widow_control = True
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        _set_run_font(lead, _docx_text_font(bold_lead), 11.5)
    run = paragraph.add_run(text)
    _set_run_font(run, _docx_text_font(text), 11.5)
    return paragraph


def _add_docx_heading(document: Document, text: str, level: int):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        _set_run_font(run, "Aptos Display" if not DEVANAGARI_PATTERN.search(text) else "Nirmala UI", 15 if level == 1 else 12)
        run.font.color.rgb = RGBColor(16, 42, 53)
    return paragraph


def _add_docx_list(document: Document, items: Iterable[str], *, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        paragraph = _add_docx_paragraph(document, item, style=style)
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)


def build_docx_lesson_notes(lesson: dict[str, Any], generated_at: str) -> bytes:
    try:
        notes = build_lesson_notes_content(lesson, generated_at)
        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)

        normal = document.styles["Normal"]
        normal.font.name = "Georgia"
        normal.font.size = Pt(11.5)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        normal_properties = normal._element.get_or_add_rPr()
        normal_fonts = normal_properties.get_or_add_rFonts()
        normal_fonts.set(qn("w:eastAsia"), "Nirmala UI")
        normal_fonts.set(qn("w:cs"), "Nirmala UI")

        brand = document.add_paragraph()
        brand.paragraph_format.space_after = Pt(9)
        brand_run = brand.add_run("Adhyantra")
        brand_run.bold = True
        brand_run.font.color.rgb = RGBColor(15, 143, 131)
        _set_run_font(brand_run, "Aptos Display", 12)

        title = document.add_paragraph(style="Title")
        title.paragraph_format.space_after = Pt(10)
        title_run = title.add_run(notes.title)
        title_run.bold = True
        title_run.font.color.rgb = RGBColor(0, 0, 0)
        _set_run_font(title_run, _docx_text_font(notes.title), 24)

        metadata = (
            ("Exam", notes.exam),
            ("Subject", notes.subject),
            ("Topic", notes.topic),
            ("Teaching approach", notes.teaching_mode),
            ("Lesson mode", notes.lesson_mode),
            ("Exported", _formatted_date(notes.generated_at)),
        )
        for label, value in metadata:
            if value:
                _add_docx_paragraph(document, value, bold_lead=f"{label}: ")

        _add_docx_heading(document, "Core explanation", 1)
        if notes.simple_explanation:
            _add_docx_paragraph(document, notes.simple_explanation, bold_lead="Key idea: ")
        if notes.detailed_explanation:
            _add_docx_paragraph(document, notes.detailed_explanation)

        for item in notes.sections:
            _add_docx_heading(document, item.title, 1)
            if item.summary:
                _add_docx_paragraph(document, item.summary)
            if item.bullets:
                _add_docx_list(document, item.bullets)
            if item.examples:
                _add_docx_heading(document, "Examples", 2)
                _add_docx_list(document, item.examples)
            if item.remember_points:
                _add_docx_heading(document, "Remember", 2)
                _add_docx_list(document, item.remember_points)
            if item.revision_cues:
                _add_docx_heading(document, "Revision cues", 2)
                _add_docx_list(document, item.revision_cues)

        if notes.questions:
            _add_docx_heading(document, "Check your understanding", 1)
            _add_docx_list(document, notes.questions, numbered=True)

        document.core_properties.title = notes.title
        document.core_properties.subject = f"{notes.exam} {notes.subject} lesson notes"
        document.core_properties.author = "Adhyantra"
        stream = BytesIO()
        document.save(stream)
        result = stream.getvalue()
        if not result.startswith(b"PK"):
            raise LessonDocumentGenerationError("The generated Word document was not a valid package.")
        return result
    except LessonDocumentGenerationError:
        raise
    except Exception as exc:
        raise LessonDocumentGenerationError("The Word lesson notes could not be generated.") from exc
